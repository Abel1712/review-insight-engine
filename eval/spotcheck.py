"""Abel's human checks, as formatted Excel sheets:
  judge: 20 LLM-judge decisions from Step 7     -> metrics.json["judge_spotcheck"]
  gold:  20 of Claude's gold_test labels         -> metrics.json["gold_spotcheck"] (human-AI agreement on the answer key)

  python eval/spotcheck.py --make judge|gold   -> eval/<kind>_spotcheck.xlsx
  python eval/spotcheck.py judge|gold          -> reads Abel's answers, records agreement
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from config import EVAL, SEED, update_metrics  # noqa: E402

XLSX = EVAL / "judge_spotcheck.xlsx"  # kept for evaluate.py, which builds the judge sheet automatically
SHEETS = {
    "judge": {"path": EVAL / "judge_spotcheck.xlsx", "answer": "Do you agree with the judge? (yes / no)",
              "instructions": "For each row: read the review, then the two labels. The AI judge decided whether they "
                              "describe the SAME issue. Write yes if you agree with the judge, no if you don't. "
                              "Save the file when done.",
              "widths": (4, 60, 32, 32, 16, 18, 24)},
    "gold": {"path": EVAL / "gold_spotcheck.xlsx", "answer": "Are these labels right? (yes / no)",
             "instructions": "For each row: read the review, then Claude's labels (every complaint or praise it found). "
                             "Write yes if the labels are correct and nothing important is missing; no if something is "
                             "wrong or missing (optionally say what in the last column). Save the file when done.",
             "widths": (4, 9, 62, 42, 20, 36)},
}


def _judge_rows() -> pd.DataFrame:
    log = pd.read_csv(EVAL / "judge_spotcheck.csv", encoding="utf-8-sig")
    test = pd.read_csv(EVAL / "gold_test.csv", encoding="utf-8-sig")
    review_for = {item.strip(): r.text for r in test.itertuples() for item in str(r.gold_complaints).split(";")}
    return pd.DataFrame({
        "#": range(1, len(log) + 1),
        "Review (what the user wrote)": log["a"].map(lambda a: review_for.get(str(a).strip(), "")),
        "Gold label (reference)": log["a"],
        "System output": log["b"],
        "Compared at": log["kind"].map({"phrase": "extracted phrase", "theme": "theme name"}),
        "Judge said: same issue?": log["judge_same"].map({True: "YES, same", False: "NO, different"}),
        SHEETS["judge"]["answer"]: "",
    })


def _gold_rows() -> pd.DataFrame:
    test = pd.read_csv(EVAL / "gold_test.csv", encoding="utf-8-sig").sample(n=20, random_state=SEED)
    return pd.DataFrame({
        "#": range(1, len(test) + 1),
        "App": test["app"].values,
        "Review (what the user wrote)": test["text"].values,
        "Claude's labels": test["gold_complaints"].str.replace(";", "\n").values,
        SHEETS["gold"]["answer"]: "",
        "What's wrong or missing? (optional)": "",
    })


def make(kind: str) -> None:
    spec = SHEETS[kind]
    sheet = _judge_rows() if kind == "judge" else _gold_rows()
    with pd.ExcelWriter(spec["path"], engine="openpyxl") as xw:
        sheet.to_excel(xw, index=False, startrow=2, sheet_name="spotcheck")
    wb = load_workbook(spec["path"])
    ws = wb["spotcheck"]
    last = chr(ord("A") + len(sheet.columns) - 1)
    ws["A1"] = spec["instructions"]
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"A1:{last}1")
    ws.row_dimensions[1].height = 34
    for col, width in zip("ABCDEFGH", spec["widths"]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[3]:
        cell.font = Font(bold=True)
    answer_col = chr(ord("A") + list(sheet.columns).index(spec["answer"]))
    for r in range(3, ws.max_row + 1):
        ws[f"{answer_col}{r}"].fill = PatternFill("solid", fgColor="FFF4CC")
    ws.freeze_panes = "A4"
    wb.save(spec["path"])
    print(f"Wrote {spec['path']} ({len(sheet)} rows)")


def record(kind: str) -> None:
    spec = SHEETS[kind]
    sheet = pd.read_excel(spec["path"], sheet_name="spotcheck", header=2)
    answers = sheet[spec["answer"]].astype(str).str.strip().str.lower()
    filled = answers.isin(["yes", "no", "y", "n"])
    if not filled.all():
        sys.exit(f"{(~filled).sum()} rows not answered yet (rows {list(sheet.loc[~filled, '#'])}). Write yes or no.")
    agree = answers.str.startswith("y")
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "analysis"))
    from trend import wilson
    lo, hi = wilson(int(agree.sum()), len(sheet))
    result = {"n": len(sheet), "n_agree": int(agree.sum()), "agreement": round(float(agree.mean()), 4),
              "agreement_ci95": [round(lo, 4), round(hi, 4)], "checked_by": "Abel (by hand)"}
    if kind == "gold":
        notes = sheet.iloc[:, -1].dropna().astype(str).str.strip()
        result["notes"] = [n for n in notes if n and n.lower() != "nan"]
    update_metrics(**{f"{kind}_spotcheck": result})
    label = "Judge-human" if kind == "judge" else "Human agreement with Claude's gold labels"
    print(f"{label}: {agree.sum()}/{len(sheet)} = {agree.mean():.0%} (95% CI {lo:.0%}-{hi:.0%})")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--make"]
    kind = args[0] if args else "judge"
    make(kind) if "--make" in sys.argv else record(kind)
